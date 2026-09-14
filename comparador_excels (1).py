# -*- coding: utf-8 -*-
"""
comparador_excels.py

Script estandarizado para comparar dos archivos Excel y encontrar las
filas que hacen "match" por nombre (o por cualquier otra columna clave
que yo defina). Lo armé para no tener que buscar manualmente nombre por
nombre entre dos listas: aquí solo indico las rutas de los dos archivos,
la hoja y la columna que quiero comparar, y el script me regresa un
Excel con 3 pestañas:
    - Coincidencias        -> filas que sí hicieron match en ambos archivos
    - Solo_en_Archivo_1    -> filas de Archivo 1 que no aparecieron en Archivo 2
    - Solo_en_Archivo_2    -> filas de Archivo 2 que no aparecieron en Archivo 1

La comparación normaliza el texto (quita espacios extra, mayúsculas/
minúsculas y acentos) para que "Juan Pérez", "juan perez " y "JUAN PEREZ"
se consideren la misma coincidencia.

Uso:
    1. Ajustar la sección "CONFIGURACIÓN" más abajo con mis rutas,
       hojas y columnas reales.
    2. Correr el script.
    3. Revisar el archivo de salida en /salida/comparacion_resultado.xlsx

Autor: Annette Michell Valenzuela Rodríguez
"""

import os
import unicodedata
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# =============================================================================
# 1. CONFIGURACIÓN
# =============================================================================
# Esta es la única sección que necesito tocar cada vez que cambio de
# archivos a comparar.

ARCHIVO_1 = "archivo_1.xlsx"     # ruta del primer Excel
HOJA_1 = 0                       # nombre o índice de la hoja (0 = primera)
COLUMNA_CLAVE_1 = "Nombre"       # columna que voy a usar para comparar

ARCHIVO_2 = "archivo_2.xlsx"     # ruta del segundo Excel
HOJA_2 = 0
COLUMNA_CLAVE_2 = "Nombre"

CARPETA_SALIDA = "salida"
ARCHIVO_SALIDA = os.path.join(CARPETA_SALIDA, "comparacion_resultado.xlsx")

FUENTE = "Arial"


# =============================================================================
# 2. NORMALIZACIÓN DE TEXTO
# =============================================================================
# Comparar strings "a lo bruto" casi nunca hace match porque vienen con
# mayúsculas distintas, espacios de más o acentos. Por eso armo una
# columna auxiliar normalizada y comparo sobre esa, sin tocar el texto
# original que se muestra en el resultado final.

def normalizar_texto(texto):
    """
    Deja el texto en minúsculas, sin espacios al inicio/final, sin
    espacios dobles y sin acentos, para que la comparación sea confiable.
    """
    if pd.isna(texto):
        return ""
    texto = str(texto).strip().lower()
    texto = " ".join(texto.split())  # colapsa espacios múltiples
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return texto


# =============================================================================
# 3. CARGA DE DATOS
# =============================================================================

def cargar_excel(ruta, hoja, columna_clave):
    """
    Carga un Excel y agrega la columna auxiliar '_clave_normalizada' que
    se usa para hacer el cruce entre los dos archivos.
    """
    df = pd.read_excel(ruta, sheet_name=hoja)

    if columna_clave not in df.columns:
        raise ValueError(
            f"La columna '{columna_clave}' no existe en '{ruta}'. "
            f"Columnas disponibles: {list(df.columns)}"
        )

    df["_clave_normalizada"] = df[columna_clave].apply(normalizar_texto)
    # Quito filas donde la clave quedó vacía (celdas en blanco), no aportan
    # nada a la comparación y solo generan falsos matches entre vacíos.
    df = df[df["_clave_normalizada"] != ""].reset_index(drop=True)
    return df


# =============================================================================
# 4. COMPARACIÓN
# =============================================================================

def comparar(df1, df2):
    """
    Compara los dos DataFrames por la columna normalizada y regresa tres
    DataFrames: coincidencias, solo en df1, solo en df2.
    """
    claves_1 = set(df1["_clave_normalizada"])
    claves_2 = set(df2["_clave_normalizada"])

    claves_comunes = claves_1 & claves_2
    claves_solo_1 = claves_1 - claves_2
    claves_solo_2 = claves_2 - claves_1

    # Para las coincidencias, uno lado a lado la fila de archivo 1 y la
    # de archivo 2 que hicieron match, para poder comparar el resto de
    # sus columnas si hace falta.
    coincidencias = pd.merge(
        df1, df2,
        on="_clave_normalizada",
        suffixes=(" (Archivo 1)", " (Archivo 2)"),
    )

    solo_1 = df1[df1["_clave_normalizada"].isin(claves_solo_1)].copy()
    solo_2 = df2[df2["_clave_normalizada"].isin(claves_solo_2)].copy()

    # Quito la columna auxiliar de los resultados finales, ya cumplió su función.
    coincidencias = coincidencias.drop(columns=["_clave_normalizada"])
    solo_1 = solo_1.drop(columns=["_clave_normalizada"])
    solo_2 = solo_2.drop(columns=["_clave_normalizada"])

    return coincidencias, solo_1, solo_2, claves_comunes


# =============================================================================
# 5. EXPORTAR RESULTADOS A EXCEL (con formato)
# =============================================================================

def escribir_hoja(ws, df, color_encabezado):
    """Escribe un DataFrame en una hoja de openpyxl con formato estándar."""
    encabezado_relleno = PatternFill(start_color=color_encabezado,
                                      end_color=color_encabezado,
                                      fill_type="solid")
    encabezado_fuente = Font(name=FUENTE, bold=True, color="FFFFFF")
    fuente_normal = Font(name=FUENTE)

    # Encabezados
    for col_idx, columna in enumerate(df.columns, start=1):
        celda = ws.cell(row=1, column=col_idx, value=str(columna))
        celda.font = encabezado_fuente
        celda.fill = encabezado_relleno
        celda.alignment = Alignment(horizontal="center")

    # Datos
    for fila_idx, fila in enumerate(df.itertuples(index=False), start=2):
        for col_idx, valor in enumerate(fila, start=1):
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            celda.font = fuente_normal

    # Ajuste de ancho de columnas según el contenido
    for col_idx, columna in enumerate(df.columns, start=1):
        largo_max = max(
            [len(str(columna))] + [len(str(v)) for v in df[columna].astype(str)]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(largo_max + 3, 45)

    ws.freeze_panes = "A2"


def exportar_resultados(coincidencias, solo_1, solo_2, ruta_salida):
    wb = Workbook()

    ws_coincidencias = wb.active
    ws_coincidencias.title = "Coincidencias"
    escribir_hoja(ws_coincidencias, coincidencias, "2E7D32")  # verde

    ws_solo_1 = wb.create_sheet("Solo_en_Archivo_1")
    escribir_hoja(ws_solo_1, solo_1, "C0392B")  # rojo

    ws_solo_2 = wb.create_sheet("Solo_en_Archivo_2")
    escribir_hoja(ws_solo_2, solo_2, "C0392B")  # rojo

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    wb.save(ruta_salida)


# =============================================================================
# 6. FLUJO PRINCIPAL
# =============================================================================

def main():
    df1 = cargar_excel(ARCHIVO_1, HOJA_1, COLUMNA_CLAVE_1)
    df2 = cargar_excel(ARCHIVO_2, HOJA_2, COLUMNA_CLAVE_2)

    coincidencias, solo_1, solo_2, claves_comunes = comparar(df1, df2)

    exportar_resultados(coincidencias, solo_1, solo_2, ARCHIVO_SALIDA)

    # Resumen rápido en consola para no tener que abrir el Excel nada
    # más para saber cuántos matches hubo.
    print("=" * 50)
    print("RESUMEN DE COMPARACIÓN")
    print("=" * 50)
    print(f"Filas en {ARCHIVO_1}: {len(df1)}")
    print(f"Filas en {ARCHIVO_2}: {len(df2)}")
    print(f"Coincidencias encontradas: {len(claves_comunes)}")
    print(f"Solo en {ARCHIVO_1}: {len(solo_1)}")
    print(f"Solo en {ARCHIVO_2}: {len(solo_2)}")
    print(f"\nResultado guardado en: {ARCHIVO_SALIDA}")


if __name__ == "__main__":
    main()
